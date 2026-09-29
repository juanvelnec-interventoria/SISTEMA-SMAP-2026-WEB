# -*- coding: utf-8 -*-
"""Actualizador local -> SISTEMA SMAP 2026 WEB."""
import json, os, math, urllib.request, urllib.error
from pathlib import Path
from datetime import datetime, date
from openpyxl import load_workbook

WEB_URL=os.environ.get('SMAP_WEB_URL','https://sistema-smap-2026-web.onrender.com').rstrip('/')
TOKEN=os.environ.get('SMAP_SYNC_TOKEN','SMAP2026_SYNC_751_VELNEC')
EXCEL=Path(os.environ.get('SMAP_EXCEL',r'C:\Users\USER\Desktop\INTERVENTORIA VELNEC\SMAP\SMAP_VELNEC_2026_JCA.xlsx'))
SHEET='Tabla SMAP-CIV'

def norm(s): return ' '.join(str(s or '').replace('\\n',' ').split()).strip().lower()
def blank(v): return v is None or (isinstance(v,float) and math.isnan(v)) or str(v).strip()=='' or str(v).strip().lower() in {'nan','#n/a','none'}
def serial(v):
    if isinstance(v,(datetime,date)): return v.isoformat()
    if isinstance(v,float) and math.isnan(v): return None
    if isinstance(v,float) and v.is_integer(): return int(v)
    return v
def idx(headers,*names):
    for name in names:
        name=norm(name)
        for i,header in enumerate(headers):
            if header==name: return i
    return -1

print('')
print('='*60)
print('       SISTEMA SMAP 2026')
print('       ACTUALIZAR TABLERO WEB')
print('='*60)
print('')
print('Excel:'); print(EXCEL); print('')
print('Servidor:'); print(WEB_URL); print('')

if not EXCEL.exists():
    print('ERROR: No se encontró el Excel maestro.'); print(EXCEL)
    input('\nPresione ENTER para cerrar...'); raise SystemExit(1)

try:
    wb=load_workbook(EXCEL,read_only=True,data_only=True)
    if SHEET not in wb.sheetnames:
        print(f"ERROR: No existe la hoja '{SHEET}'.")
        print('Hojas disponibles:',wb.sheetnames); wb.close()
        input('\nPresione ENTER para cerrar...'); raise SystemExit(1)
    ws=wb[SHEET]; it=ws.iter_rows(min_row=2,values_only=True); headers=list(next(it)); H=[norm(x) for x in headers]
    I={
      'consecutivo':idx(H,'Consecutivo'),'mes':idx(H,'Mes'),'anio':idx(H,'Año'),'solicitante':idx(H,'Solicitante'),
      'direccion':idx(H,'Dirección'),'lat':idx(H,'Coordenadas Latitud'),'lon':idx(H,'Coordenadas Longitud'),
      'barrio':idx(H,'Barrio/Vereda'),'comuna':idx(H,'Comuna'),'nombreComuna':idx(H,'Nombre comuna'),
      'falla':idx(H,'N° Falla'),'tipoRecorrido':idx(H,'TIPO RECORRIDO'),'oficio':idx(H,'Oficio de envio'),
      'fechaEnvio':idx(H,'Fecha de envio CIV'),'estado':idx(H,'Estado de SMAP'),'noSS':idx(H,'No. de SS'),
      'tipoSS':idx(H,'Tipo de SS'),'fechaSS':idx(H,'Fecha recibo del numero de SS'),
      'diasSS':idx(H,'Días Trascurridos entrega numero de SS'),'cumpleSS':idx(H,'Cumple ANS entrega de numero de SS'),
      'nivel':idx(H,'Tipo de Nivel'),'fechaEntrega':idx(H,'Fecha entrega de Respuesta'),
      'diasRespuesta':idx(H,'Días Trascurridos entrega de Respuesta'),'vencimiento':idx(H,'VENCIMIENTO ENTREGA RESPUESTA'),
      'cumpleRespuesta':idx(H,'Cumple ANS entrega de Respuesta'),'fechaRespuesta':idx(H,'Fecha de Respuesta'),
      'radicadoResp':idx(H,'Radicado Respuesta EPM2'),'observaciones':idx(H,'Observaciones'),
      'aval':idx(H,'Radicado AVAL VELNEC'),'fechaAval':idx(H,'Fecha de AVAL'),'aprobDistrito':idx(H,'RAD.APROBACION DISTRITO'),
      'fechaAprob':idx(H,'Fecha de APROB DISTRITO'),'ssRevision':idx(H,'FECHA RESPUEST SS REVISAR_EJECUCION'),
      'diasRevision':idx(H,'Días Trascurridos entrega de SS REVISAR'),'vencRevision':idx(H,'VENCIMIENTO'),
      'cumpleRevision':idx(H,'Cumple ANS entrega de SS REVISAR'),'radEjec':idx(H,'Radicado EJECUCION VELNEC'),
      'fechaEjec':idx(H,'Fecha de RADI EJECUCION'),'iniciales':idx(H,'Iniciales'),'interventor':idx(H,'INTERVENTOR'),
      'tecnico':idx(H,'Tecnicos'),'fallaSmap':idx(H,'FALLAS smap')}
    exact=next((i for i,h in enumerate(H) if h=='interventor'),-1)
    I['interventor']=exact if exact>=0 else 5
    data=[]
    for row in it:
        if not row or I['consecutivo']<0: continue
        con=row[I['consecutivo']] if I['consecutivo']<len(row) else None
        if blank(con): continue
        item={k:(serial(row[j]) if 0<=j<len(row) else None) for k,j in I.items()}
        item['pendienteActualizacion']=not bool(item.get('anio') or item.get('mes') or item.get('direccion'))
        data.append(item)
    wb.close()
except Exception as e:
    print(''); print('='*60); print('ERROR LEYENDO EL EXCEL'); print('='*60); print(''); print(e); print('')
    input('Presione ENTER para cerrar...'); raise SystemExit(1)

print('Excel leído correctamente.'); print(f'Registros encontrados: {len(data)}'); print('')
payload=json.dumps({'data':data},ensure_ascii=False,separators=(',',':')).encode('utf-8')
url=WEB_URL+'/api/sync'
print('Conectando con:'); print(url); print(''); print('Enviando información...'); print('')
request=urllib.request.Request(url,data=payload,method='POST',headers={'Content-Type':'application/json','X-SMAP-TOKEN':TOKEN,'User-Agent':'SISTEMA-SMAP-2026'})
try:
    with urllib.request.urlopen(request,timeout=60) as response:
        result=json.loads(response.read().decode('utf-8'))
    if not result.get('ok'):
        print('ERROR: El servidor rechazó la actualización.'); print(result)
        input('\nPresione ENTER para cerrar...'); raise SystemExit(1)
    print(''); print('='*60); print('       ACTUALIZACIÓN COMPLETADA'); print('='*60); print('')
    print(f"Registros enviados: {result.get('rows')}")
    print(f"Fecha actualización: {result.get('updated_at')}")
    print(''); print('El tablero web ya recibió la información.'); print(''); print('URL:',WEB_URL); print(''); print('='*60)
except urllib.error.HTTPError as e:
    try: body=e.read().decode('utf-8')
    except Exception: body=''
    print(''); print('='*60); print('ERROR HTTP AL ACTUALIZAR'); print('='*60); print('')
    print('Código:',e.code); print('URL:',url); print('Respuesta:',body); print('')
    if e.code==401: print('El servidor rechazó el TOKEN. Verifique que server.py en Render use el mismo token.')
    elif e.code==404: print('La ruta /api/sync no fue encontrada.')
    elif e.code==500: print('Render encontró un error interno.')
    print('')
except Exception as e:
    print(''); print('='*60); print('ERROR DE CONEXIÓN'); print('='*60); print(''); print(e); print('')
input('Presione ENTER para cerrar...')
