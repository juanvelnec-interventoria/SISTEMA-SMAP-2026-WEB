# -*- coding: utf-8 -*-
"""
SISTEMA SMAP 2026 - SERVIDOR LOCAL + RENDER

Este servidor funciona en dos modos:

1. LOCAL:
   - Lee directamente el Excel maestro.
   - Sirve el tablero local.
   - Actualiza los datos cuando cambia el Excel.

2. RENDER:
   - No necesita el Excel local.
   - Recibe los datos enviados por actualizar_smap_web.py
     mediante POST /api/sync.
   - Guarda los datos en data.json.
   - Sirve esos datos al tablero web.

IMPORTANTE:
El token de sincronización debe coincidir con el usado por
actualizar_smap_web.py.
"""

from http.server import ThreadingHTTPServer, SimpleHTTPRequestHandler
from pathlib import Path
from urllib.parse import urlparse
from datetime import datetime, date
import json
import math
import os
import threading

try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None


# ============================================================
# CONFIGURACIÓN GENERAL
# ============================================================

ROOT = Path(__file__).resolve().parent

CFG = ROOT / "config.json"
DATA_FILE = ROOT / "data.json"
DEMO_FILE = ROOT / "demo_data.json"

try:
    config = json.loads(CFG.read_text(encoding="utf-8"))
except Exception:
    config = {}

CONFIG_EXCEL = Path(
    config.get(
        "excel_path",
        r"C:\Users\USER\Desktop\INTERVENTORIA VELNEC\SMAP\SMAP_VELNEC_2026_JCA.xlsx"
    )
)

SHEET = config.get(
    "sheet_name",
    "Tabla SMAP-CIV"
)


# ============================================================
# PUERTO
# ============================================================
# Render proporciona PORT automáticamente.
# En local, si no existe PORT, usamos 8876.

PORT = int(
    os.environ.get(
        "PORT",
        "8876"
    )
)


# ============================================================
# TOKEN DE SINCRONIZACIÓN
# ============================================================
# Debe coincidir con actualizar_smap_web.py.
#
# Para una instalación más segura en el futuro se puede definir
# SMAP_SYNC_TOKEN como variable de entorno en Render.
#
# Mientras tanto, este valor permite trabajar sin depender de
# Environment Variables de Render.

SYNC_TOKEN = os.environ.get(
    "SMAP_SYNC_TOKEN",
    "SMAP2026_SYNC_751_VELNEC"
)


# ============================================================
# ESTADO / CACHE
# ============================================================

LOCK = threading.Lock()

CACHE = {
    "mtime": None,
    "data": [],
    "error": None,
    "loaded_at": None,
    "refreshing": False
}


# ============================================================
# FUNCIONES GENERALES
# ============================================================

def norm(s):
    return " ".join(
        str(s or "")
        .replace("\n", " ")
        .split()
    ).strip().lower()


def blank(v):
    return (
        v is None
        or (
            isinstance(v, float)
            and math.isnan(v)
        )
        or str(v).strip() == ""
        or str(v).strip().lower()
        in {
            "nan",
            "#n/a",
            "none"
        }
    )


def clean(v):
    return "" if blank(v) else str(v).strip()


def serial(v):

    if v is None:
        return None

    if isinstance(v, (datetime, date)):
        return v.isoformat()

    if isinstance(v, float) and math.isnan(v):
        return None

    if isinstance(v, float) and v.is_integer():
        return int(v)

    return v


# ============================================================
# UBICAR EXCEL
# ============================================================

def resolve_excel_path():

    if CONFIG_EXCEL.exists():
        return CONFIG_EXCEL

    candidates = []

    home = Path.home()

    roots = [
        home / "Desktop" / "INTERVENTORIA VELNEC" / "SMAP",
        home / "OneDrive" / "Desktop" / "INTERVENTORIA VELNEC" / "SMAP",

        Path(
            r"C:\Users\USER\Desktop\INTERVENTORIA VELNEC\SMAP"
        ),

        Path(
            r"C:\Users\USER\OneDrive\Desktop\INTERVENTORIA VELNEC\SMAP"
        )
    ]

    for root in roots:

        try:

            if root.exists():

                candidates += list(
                    root.glob(
                        "SMAP_VELNEC_2026_JCA*.xlsx"
                    )
                )

        except Exception:
            pass

    if candidates:

        return max(
            candidates,
            key=lambda x: x.stat().st_mtime
        )

    return CONFIG_EXCEL


EXCEL = resolve_excel_path()


# ============================================================
# LEER EXCEL
# ============================================================

def read_excel():

    global EXCEL

    EXCEL = resolve_excel_path()

    if load_workbook is None:

        raise RuntimeError(
            "Falta openpyxl. Instale Python y openpyxl."
        )

    if not EXCEL.exists():

        raise FileNotFoundError(
            f"No se encontró el Excel maestro: {EXCEL}"
        )

    mtime = EXCEL.stat().st_mtime

    wb = load_workbook(
        EXCEL,
        read_only=True,
        data_only=True
    )

    if SHEET not in wb.sheetnames:

        wb.close()

        raise KeyError(
            f"No existe la hoja {SHEET!r}. "
            f"Hojas disponibles: {wb.sheetnames}"
        )

    ws = wb[SHEET]

    it = ws.iter_rows(
        min_row=2,
        values_only=True
    )

    headers = list(next(it))

    H = [
        norm(x)
        for x in headers
    ]


    def idx(*names):

        ns = [
            norm(x)
            for x in names
        ]

        for n in ns:

            for i, h in enumerate(H):

                if h == n:
                    return i

        return -1


    # ========================================================
    # COLUMNAS
    # ========================================================

    I = {

        "consecutivo":
            idx("Consecutivo"),

        "mes":
            idx("Mes"),

        "anio":
            idx("Año"),

        "solicitante":
            idx("Solicitante"),

        "direccion":
            idx("Dirección"),

        "lat":
            idx("Coordenadas Latitud"),

        "lon":
            idx("Coordenadas Longitud"),

        "barrio":
            idx("Barrio/Vereda"),

        "comuna":
            idx("Comuna"),

        "nombreComuna":
            idx("Nombre comuna"),

        "falla":
            idx("N° Falla"),

        "tipoRecorrido":
            idx("TIPO RECORRIDO"),

        "oficio":
            idx("Oficio de envio"),

        "fechaEnvio":
            idx("Fecha de envio CIV"),

        "estado":
            idx("Estado de SMAP"),

        "noSS":
            idx("No. de SS"),

        "tipoSS":
            idx("Tipo de SS"),

        "fechaSS":
            idx("Fecha recibo del numero de SS"),

        "diasSS":
            idx(
                "Días Trascurridos entrega numero de SS"
            ),

        "cumpleSS":
            idx(
                "Cumple ANS entrega de numero de SS"
            ),

        "nivel":
            idx("Tipo de Nivel"),

        "fechaEntrega":
            idx(
                "Fecha entrega de Respuesta"
            ),

        "diasRespuesta":
            idx(
                "Días Trascurridos entrega de Respuesta"
            ),

        "vencimiento":
            idx(
                "VENCIMIENTO ENTREGA RESPUESTA"
            ),

        "cumpleRespuesta":
            idx(
                "Cumple ANS entrega de Respuesta"
            ),

        "fechaRespuesta":
            idx(
                "Fecha de Respuesta"
            ),

        "radicadoResp":
            idx(
                "Radicado Respuesta EPM2"
            ),

        "observaciones":
            idx("Observaciones"),

        "aval":
            idx(
                "Radicado AVAL VELNEC"
            ),

        "fechaAval":
            idx(
                "Fecha de AVAL"
            ),

        "aprobDistrito":
            idx(
                "RAD.APROBACION DISTRITO"
            ),

        "fechaAprob":
            idx(
                "Fecha de APROB DISTRITO"
            ),

        "ssRevision":
            idx(
                "FECHA RESPUEST SS REVISAR_EJECUCION"
            ),

        "diasRevision":
            idx(
                "Días Trascurridos entrega de SS REVISAR"
            ),

        "vencRevision":
            idx("VENCIMIENTO"),

        "cumpleRevision":
            idx(
                "Cumple ANS entrega de SS REVISAR"
            ),

        "radEjec":
            idx(
                "Radicado EJECUCION VELNEC"
            ),

        "fechaEjec":
            idx(
                "Fecha de RADI EJECUCION"
            ),

        "iniciales":
            idx("Iniciales"),

        "interventor":
            idx("INTERVENTOR"),

        "tecnico":
            idx("Tecnicos"),

        "fallaSmap":
            idx("FALLAS smap")
    }


    # ========================================================
    # INTERVENTOR REAL
    # ========================================================
    # En el Excel actual la asignación por solicitud está en F.
    # Evitamos confundirla con N°INTERVENTOR.

    exact_interventor = next(
        (
            i
            for i, h in enumerate(H)
            if h == "interventor"
        ),
        -1
    )

    if exact_interventor >= 0:

        I["interventor"] = exact_interventor

    else:

        candidate = 5

        if candidate < len(headers):

            I["interventor"] = candidate


    # ========================================================
    # CONSTRUIR DATOS
    # ========================================================

    data = []

    for row in it:

        if not row:
            continue

        if I["consecutivo"] < 0:
            continue

        con = (
            row[I["consecutivo"]]
            if I["consecutivo"] < len(row)
            else None
        )

        if blank(con):
            continue

        d = {}

        for k, j in I.items():

            if (
                j >= 0
                and j < len(row)
            ):

                d[k] = serial(
                    row[j]
                )

            else:

                d[k] = None


        # ====================================================
        # PENDIENTE DE ACTUALIZACIÓN
        # ====================================================

        d["pendienteActualizacion"] = not bool(
            d.get("anio")
            or d.get("mes")
            or d.get("direccion")
        )

        data.append(d)


    wb.close()

    return data, mtime


# ============================================================
# GUARDAR DATA RECIBIDA
# ============================================================

def save_data(data):

    payload = {
        "updated_at": datetime.now().isoformat(
            timespec="seconds"
        ),
        "rows": len(data),
        "data": data
    }

    DATA_FILE.write_text(
        json.dumps(
            payload,
            ensure_ascii=False,
            separators=(",", ":")
        ),
        encoding="utf-8"
    )


# ============================================================
# CARGAR DATA GUARDADA
# ============================================================

def load_saved_data():

    try:

        if not DATA_FILE.exists():
            return []

        obj = json.loads(
            DATA_FILE.read_text(
                encoding="utf-8"
            )
        )

        if isinstance(obj, dict):

            data = obj.get(
                "data",
                []
            )

        else:

            data = obj

        if isinstance(data, list):

            return data

    except Exception as e:

        with LOCK:
            CACHE["error"] = str(e)

    return []


# ============================================================
# DEMO
# ============================================================

def load_demo():

    try:

        if DEMO_FILE.exists():

            obj = json.loads(
                DEMO_FILE.read_text(
                    encoding="utf-8"
                )
            )

            if isinstance(obj, dict):

                obj = obj.get(
                    "data",
                    []
                )

            if isinstance(obj, list):

                with LOCK:

                    CACHE.update(
                        data=obj,
                        error=None,
                        loaded_at="datos demo",
                        refreshing=False
                    )

                return obj

    except Exception as e:

        with LOCK:
            CACHE["error"] = (
                f"Demo: {e}"
            )

    return []


# ============================================================
# ACTUALIZAR CACHE DESDE EXCEL
# ============================================================

def refresh_worker():

    try:

        data, mtime = read_excel()

        with LOCK:

            CACHE.update(
                mtime=mtime,
                data=data,
                error=None,
                loaded_at=datetime.now().isoformat(
                    timespec="seconds"
                ),
                refreshing=False
            )

    except Exception as e:

        with LOCK:

            CACHE.update(
                error=str(e),
                refreshing=False
            )


def start_refresh(force=False):

    global EXCEL

    EXCEL = resolve_excel_path()

    with LOCK:

        if CACHE.get("refreshing"):
            return False

        current = (
            EXCEL.stat().st_mtime
            if EXCEL.exists()
            else None
        )

        if (
            not force
            and current is not None
            and CACHE.get("mtime") == current
        ):

            return False

        CACHE["refreshing"] = True

    threading.Thread(
        target=refresh_worker,
        daemon=True
    ).start()

    return True


# ============================================================
# CARGAR DATOS
# ============================================================

def load_excel():

    global EXCEL

    EXCEL = resolve_excel_path()

    with LOCK:

        data = CACHE.get(
            "data"
        ) or []

    # Si hay Excel local, usamos el Excel.
    if EXCEL.exists():

        if data:

            start_refresh(False)
            return data

        data, mtime = read_excel()

        with LOCK:

            CACHE.update(
                mtime=mtime,
                data=data,
                error=None,
                loaded_at=datetime.now().isoformat(
                    timespec="seconds"
                ),
                refreshing=False
            )

        return data

    # Si estamos en Render, usamos data.json.
    saved = load_saved_data()

    if saved:

        with LOCK:

            CACHE.update(
                data=saved,
                error=None,
                loaded_at=datetime.now().isoformat(
                    timespec="seconds"
                ),
                refreshing=False
            )

        return saved

    # Último recurso: demo.
    return load_demo()


# ============================================================
# HANDLER HTTP
# ============================================================

class Handler(SimpleHTTPRequestHandler):


    def _send_json(
        self,
        obj,
        status=200
    ):

        raw = json.dumps(
            obj,
            ensure_ascii=False,
            separators=(",", ":")
        ).encode("utf-8")

        self.send_response(status)

        self.send_header(
            "Content-Type",
            "application/json; charset=utf-8"
        )

        self.send_header(
            "Content-Length",
            str(len(raw))
        )

        self.send_header(
            "Cache-Control",
            "no-store"
        )

        self.send_header(
            "Access-Control-Allow-Origin",
            "*"
        )

        self.end_headers()

        self.wfile.write(raw)


    def do_OPTIONS(self):

        self.send_response(204)

        self.send_header(
            "Access-Control-Allow-Origin",
            "*"
        )

        self.send_header(
            "Access-Control-Allow-Methods",
            "GET, POST, OPTIONS"
        )

        self.send_header(
            "Access-Control-Allow-Headers",
            "Content-Type, X-SMAP-TOKEN"
        )

        self.end_headers()


    def do_GET(self):

        p = urlparse(
            self.path
        ).path

        try:

            # =================================================
            # HEALTH
            # =================================================

            if p == "/api/health":

                self._send_json(
                    {
                        "ok": True,
                        "service":
                            "SISTEMA-SMAP-2026-WEB",
                        "time":
                            datetime.now().isoformat(
                                timespec="seconds"
                            ),
                        "rows":
                            len(
                                CACHE.get("data")
                                or []
                            )
                    }
                )

                return


            # =================================================
            # META
            # =================================================

            if p == "/api/meta":

                exists = EXCEL.exists()

                mt = (
                    EXCEL.stat().st_mtime
                    if exists
                    else None
                )

                self._send_json(
                    {
                        "excel_path":
                            str(EXCEL),

                        "exists":
                            exists,

                        "mtime":
                            mt,

                        "cache_mtime":
                            CACHE.get("mtime"),

                        "loaded_at":
                            CACHE.get(
                                "loaded_at"
                            ),

                        "rows":
                            len(
                                CACHE.get(
                                    "data"
                                ) or []
                            ),

                        "error":
                            CACHE.get(
                                "error"
                            ),

                        "refreshing":
                            CACHE.get(
                                "refreshing",
                                False
                            ),

                        "server_mode":
                            (
                                "LOCAL"
                                if exists
                                else "RENDER"
                            )
                    }
                )

                return


            # =================================================
            # DATA
            # =================================================

            if p == "/api/data":

                data = load_excel()

                self._send_json(
                    {
                        "data": data,

                        "meta": {

                            "excel_path":
                                str(EXCEL),

                            "mtime":
                                (
                                    EXCEL.stat().st_mtime
                                    if EXCEL.exists()
                                    else None
                                ),

                            "cache_mtime":
                                CACHE.get(
                                    "mtime"
                                ),

                            "loaded_at":
                                CACHE.get(
                                    "loaded_at"
                                ),

                            "rows":
                                len(data),

                            "refreshing":
                                CACHE.get(
                                    "refreshing",
                                    False
                                )
                        }
                    }
                )

                return


            # =================================================
            # REFRESH
            # =================================================

            if p == "/api/refresh":

                started = start_refresh(
                    True
                )

                self._send_json(
                    {
                        "ok": True,
                        "refreshing": started,
                        "rows":
                            len(
                                CACHE.get(
                                    "data"
                                ) or []
                            ),
                        "loaded_at":
                            CACHE.get(
                                "loaded_at"
                            )
                    }
                )

                return


        except Exception as e:

            with LOCK:
                CACHE["error"] = str(e)

            self._send_json(
                {
                    "ok": False,
                    "error": str(e)
                },
                500
            )

            return


        return super().do_GET()


    # ========================================================
    # POST /api/sync
    # ========================================================

    def do_POST(self):

        p = urlparse(
            self.path
        ).path

        if p != "/api/sync":

            self._send_json(
                {
                    "ok": False,
                    "error": "Ruta no encontrada"
                },
                404
            )

            return


        # ====================================================
        # VALIDAR TOKEN
        # ====================================================

        received_token = self.headers.get(
            "X-SMAP-TOKEN",
            ""
        )

        if received_token != SYNC_TOKEN:

            self._send_json(
                {
                    "ok": False,
                    "error": "No autorizado"
                },
                401
            )

            return


        # ====================================================
        # LEER CUERPO
        # ====================================================

        try:

            content_length = int(
                self.headers.get(
                    "Content-Length",
                    "0"
                )
            )

            if content_length <= 0:

                self._send_json(
                    {
                        "ok": False,
                        "error":
                            "Solicitud sin datos"
                    },
                    400
                )

                return


            raw = self.rfile.read(
                content_length
            )

            payload = json.loads(
                raw.decode("utf-8")
            )


            if isinstance(
                payload,
                dict
            ):

                data = payload.get(
                    "data",
                    []
                )

            else:

                data = payload


            if not isinstance(
                data,
                list
            ):

                self._send_json(
                    {
                        "ok": False,
                        "error":
                            "El campo data debe ser una lista"
                    },
                    400
                )

                return


            # =================================================
            # GUARDAR
            # =================================================

            save_data(data)


            with LOCK:

                CACHE.update(
                    data=data,
                    error=None,
                    loaded_at=datetime.now().isoformat(
                        timespec="seconds"
                    ),
                    refreshing=False
                )


            self._send_json(
                {
                    "ok": True,
                    "rows": len(data),
                    "updated_at":
                        CACHE.get(
                            "loaded_at"
                        ),
                    "message":
                        "Datos SMAP actualizados correctamente"
                }
            )


        except Exception as e:

            with LOCK:
                CACHE["error"] = str(e)

            self._send_json(
                {
                    "ok": False,
                    "error": str(e)
                },
                500
            )


# ============================================================
# INICIO DEL SERVIDOR
# ============================================================

if __name__ == "__main__":

    os.chdir(ROOT)

    print("")
    print("=" * 60)
    print("       SISTEMA SMAP 2026")
    print("       SERVIDOR WEB")
    print("=" * 60)
    print("")

    print(
        "Puerto:",
        PORT
    )

    print(
        "Excel:",
        EXCEL
    )

    print(
        "Hoja:",
        SHEET
    )

    print("")

    # ========================================================
    # LOCAL
    # ========================================================

    if EXCEL.exists():

        print(
            "Modo: LOCAL - Excel maestro detectado"
        )

        try:

            data, mtime = read_excel()

            with LOCK:

                CACHE.update(
                    mtime=mtime,
                    data=data,
                    error=None,
                    loaded_at=datetime.now().isoformat(
                        timespec="seconds"
                    ),
                    refreshing=False
                )

            print(
                f"Registros cargados: {len(data)}"
            )

        except Exception as e:

            print(
                "Error leyendo Excel:",
                e
            )

            load_demo()

    # ========================================================
    # RENDER
    # ========================================================

    else:

        print(
            "Modo: RENDER - Excel local no encontrado"
        )

        saved = load_saved_data()

        if saved:

            with LOCK:

                CACHE.update(
                    data=saved,
                    error=None,
                    loaded_at=datetime.now().isoformat(
                        timespec="seconds"
                    ),
                    refreshing=False
                )

            print(
                f"Registros guardados cargados: {len(saved)}"
            )

        else:

            demo = load_demo()

            print(
                f"Registros demo cargados: {len(demo)}"
            )


    print("")

    print(
        "Endpoints disponibles:"
    )

    print(
        "  /api/health"
    )

    print(
        "  /api/data"
    )

    print(
        "  /api/meta"
    )

    print(
        "  /api/refresh"
    )

    print(
        "  POST /api/sync"
    )

    print("")

    print(
        "Servidor iniciado..."
    )

    print("")


    server = ThreadingHTTPServer(
        (
            "0.0.0.0",
            PORT
        ),
        Handler
    )

    server.serve_forever()
